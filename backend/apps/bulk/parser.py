"""
Read a merchant's payout spreadsheet into rows of plain strings.

Accepts .csv (UTF-8, with or without BOM; Excel's Latin-1 as a fallback) and .xlsx.
Headers are matched loosely — "Phone", "phone_number" and "Account" all mean the
account column — so a file exported from any payroll tool usually just works.

Safety: size and row caps; the workbook is opened read-only with formulas replaced by
their cached values (openpyxl never executes them); nothing is written to disk.
"""

from __future__ import annotations

import csv
import io
import re

from .exceptions import ParseError

MAX_BYTES = 2 * 1024 * 1024       # 2 MB
MAX_ROWS = 5_000

# canonical column -> accepted header spellings (lower-case, non-alphanumerics removed)
COLUMN_ALIASES = {
    "name": {"name", "recipient", "recipientname", "employee", "employeename", "staff",
             "staffname", "fullname", "beneficiary", "payee"},
    "destination_type": {"destinationtype", "destination", "type", "method", "payto",
                         "channel", "paymenttype"},
    "institution": {"institution", "network", "bank", "bankcode", "provider", "operator",
                    "scheme", "wallet", "walletprovider"},
    "account": {"account", "accountnumber", "accountno", "number", "phone", "phonenumber",
                "msisdn", "mobile", "mobilenumber", "walletnumber"},
    "amount": {"amount", "amountghs", "ghs", "value", "salary", "pay", "netpay", "total"},
    "narrative": {"narrative", "note", "notes", "description", "reference", "memo",
                  "remark", "remarks", "purpose"},
}
REQUIRED = ("name", "account", "amount")


def _norm_header(value) -> str:
    return re.sub(r"[^a-z0-9]", "", str(value or "").strip().lower())


def _map_headers(headers: list) -> dict[int, str]:
    """Column index -> canonical name. Unknown columns are ignored."""
    mapping: dict[int, str] = {}
    seen: set[str] = set()
    for idx, raw in enumerate(headers):
        key = _norm_header(raw)
        for canonical, aliases in COLUMN_ALIASES.items():
            if key in aliases and canonical not in seen:
                mapping[idx] = canonical
                seen.add(canonical)
                break
    missing = [c for c in REQUIRED if c not in seen]
    if missing:
        raise ParseError(
            "Missing column(s): " + ", ".join(missing) +
            ". Download the sample file to see the expected layout."
        )
    return mapping


def _rows_from_table(table: list[list]) -> list[dict]:
    # First non-empty line is the header.
    start = next((i for i, r in enumerate(table) if any(str(c or "").strip() for c in r)), None)
    if start is None:
        raise ParseError("The file is empty.")
    mapping = _map_headers(table[start])
    rows: list[dict] = []
    for line_no, raw in enumerate(table[start + 1:], start=1):
        if not any(str(c or "").strip() for c in raw):
            continue                                        # skip blank lines
        row = {name: "" for name in COLUMN_ALIASES}
        for idx, canonical in mapping.items():
            if idx < len(raw) and raw[idx] is not None:
                row[canonical] = str(raw[idx]).strip()
        row["_row"] = line_no
        rows.append(row)
        if len(rows) > MAX_ROWS:
            raise ParseError(f"Too many rows — the limit is {MAX_ROWS:,} per batch.")
    if not rows:
        raise ParseError("The file has a header but no payment rows.")
    return rows


def parse(filename: str, data: bytes) -> list[dict]:
    if len(data) > MAX_BYTES:
        raise ParseError("File is larger than 2 MB. Split it into smaller batches.")
    name = (filename or "").lower()
    if name.endswith(".csv") or name.endswith(".txt"):
        return _rows_from_table(_read_csv(data))
    if name.endswith(".xlsx") or name.endswith(".xlsm"):
        return _rows_from_table(_read_xlsx(data))
    if name.endswith(".xls"):
        raise ParseError("Old Excel (.xls) isn't supported. Save the file as .xlsx or .csv.")
    raise ParseError("Upload a .csv or .xlsx file.")


def _read_csv(data: bytes) -> list[list]:
    try:
        text = data.decode("utf-8-sig")
    except UnicodeDecodeError:
        text = data.decode("latin-1")
    sample = text[:4096]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t")
    except csv.Error:
        dialect = csv.excel
    return [row for row in csv.reader(io.StringIO(text), dialect)]


def _read_xlsx(data: bytes) -> list[list]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:  # pragma: no cover
        raise ParseError("Excel support is not installed on the server.") from exc
    try:
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as exc:
        raise ParseError("Could not open this Excel file. Is it a valid .xlsx?") from exc
    try:
        ws = wb.worksheets[0]
        table = []
        for i, row in enumerate(ws.iter_rows(values_only=True)):
            if i > MAX_ROWS + 5:
                raise ParseError(f"Too many rows — the limit is {MAX_ROWS:,} per batch.")
            table.append(list(row))
        return table
    finally:
        wb.close()
